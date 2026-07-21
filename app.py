import os
import re
import openpyxl
from flask import Flask, request, render_template_string, send_file
import pytesseract
from PIL import Image

app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = '/tmp/uploads'
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)

DOCUMENT_PATTERNS = {
    'CERTIFICADO': [r'CERTIFICADO DE INSPECCI[OÓ]N T[EÉ]CNICA VEHICULAR', r'RESULTADO DE LA INSPECCI[OÓ]N.*APROBADO'],
    'INFORME': [r'INFORME DE INSPECCI[OÓ]N T[EÉ]CNICA VEHICULAR', r'DATOS DE LOS EQUIPOS'],
    'TARJETA_FISICA': [r'TARJETA DE IDENTIFICACI[OÓ]N VEHICULAR', r'SUNARP'],
    'TARJETA_VIRTUAL': [r'TARJETA DE IDENTIFICACI[OÓ]N VEHICULAR ELECTR[OÓ]NICA', r'TIVE', r'C[OÓ]DIGO DE VERIFICACI[OÓ]N'],
    'SOAT': [r'SOAT', r'AFOCAT', r'AUTOSEGURO', r'CERTIFICADO CONTRA ACCIDENTES'],
    'LICENCIA': [r'LICENCIA DE CONDUCIR', r'MINISTERIO DE TRANSPORTES Y COMUNICACIONES', r'MTC'],
    'DNI': [r'DOCUMENTO NACIONAL DE IDENTIDAD', r'RENIEC', r'CARNET DE EXTRANJERIA', r'PASAPORTE'],
    'LUNAS': [r'LUNAS OSCURECIDAS', r'AUTORIZACI[OÓ]N DE USO DE LUNAS', r'POLIC[IÍ]A NACIONAL'],
    'CERT_GLP': [r'COMBUSTION DE GLP', r'CERTIFICADO DE CONFORMIDAD DEL VEHICULO CON COMBUSTION DE GLP'],
    'CERT_GNV': [r'VEH[IÍ]CULO A GNV', r'GAS NATURAL VEHICULAR', r'CERTIFICADO DE INSPECCI[OÓ]N ANUAL DEL VEH[IÍ]CULO A GNV'],
    'VOUCHER': [r'BOLETA DE VENTA ELECTR[OÓ]NICA', r'TICKET DE RECAUDACION', r'PROX\. REV\. ANUAL'],
    'INFOGAS': [r'INFOGAS', r'infogas\.com\.pe', r'Habilitado para consumir'],
    'TUC': [r'TARJETA [UÚ]NICA DE CIRCULACI[OÓ]N', r'TUC', r'AUTORIDAD DE TRANSPORTE URBANO', r'ATU'],
    'RTV_ANTERIOR': [r'FECHA PR[OÓ]XIMA INSPECCI[OÓ]N'],
    'CONSULTA_CITV': [r'Consulta de los Certificados de Inspecci[oó]n T[eé]cnica Vehicular'],
    'GLP_INICIAL_RENO': [r'CERTIFICADO DE INSPECCI[OÓ]N DE VEH[IÍ]CULO A GLP']
}

def extract_text(image_path):
    try:
        img = Image.open(image_path)
        return pytesseract.image_to_string(img, lang='spa')
    except Exception as e:
        print(f"Error OCR en {image_path}: {e}")
        return ""

def extract_placa(text):
    match = re.search(r'\b[A-Z0-9]{3}[- ]?[A-Z0-9]{3}\b', text.upper())
    return match.group(0).replace('-', '').replace(' ', '') if match else None

def extract_correlativo(text):
    match = re.search(r'SD-\d{3}-\d{7}', text)
    return match.group(0) if match else None

def process_expediente_excel(excel_path, docs_folder):
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active

    placa_to_row = {}
    for row in range(2, ws.max_row + 1):
        placa_val = str(ws.cell(row=row, column=4).value or '').strip().replace('-', '')
        if placa_val:
            placa_to_row[placa_val] = row

    for root, dirs, files in os.walk(docs_folder):
        for fname in files:
            if not fname.lower().endswith(('.png', '.jpg', '.jpeg', '.tif', '.pdf')):
                continue

            fpath = os.path.join(root, fname)
            text = extract_text(fpath)
            placa = extract_placa(text)

            if not placa or placa not in placa_to_row:
                continue

            row = placa_to_row[placa]

            if re.search(DOCUMENT_PATTERNS['CERTIFICADO'][0], text, re.IGNORECASE):
                correlativo = extract_correlativo(text)
                if correlativo:
                    ws.cell(row=row, column=8, value=correlativo) # Col H
                    ws.cell(row=row, column=27, value="A")         # Col AA (Aprobado)
                ws.cell(row=row, column=9, value="X")             # Col I

            if re.search(DOCUMENT_PATTERNS['INFORME'][0], text, re.IGNORECASE):
                ws.cell(row=row, column=10, value="X")            # Col J

            if re.search(DOCUMENT_PATTERNS['TARJETA_VIRTUAL'][0], text, re.IGNORECASE):
                ws.cell(row=row, column=12, value="X")            # Col L
            elif re.search(DOCUMENT_PATTERNS['TARJETA_FISICA'][0], text, re.IGNORECASE):
                ws.cell(row=row, column=11, value="X")            # Col K

            if any(re.search(p, text, re.IGNORECASE) for p in DOCUMENT_PATTERNS['SOAT']):
                ws.cell(row=row, column=13, value="X")            # Col M

            if any(re.search(p, text, re.IGNORECASE) for p in DOCUMENT_PATTERNS['LICENCIA']):
                ws.cell(row=row, column=14, value="X")            # Col N

            if any(re.search(p, text, re.IGNORECASE) for p in DOCUMENT_PATTERNS['DNI']):
                ws.cell(row=row, column=15, value="X")            # Col O

            if any(re.search(p, text, re.IGNORECASE) for p in DOCUMENT_PATTERNS['LUNAS']):
                ws.cell(row=row, column=16, value="X")            # Col P

            if any(re.search(p, text, re.IGNORECASE) for p in DOCUMENT_PATTERNS['CERT_GLP']):
                ws.cell(row=row, column=17, value="X")            # Col Q

            if any(re.search(p, text, re.IGNORECASE) for p in DOCUMENT_PATTERNS['CERT_GNV']):
                ws.cell(row=row, column=18, value="X")            # Col R

            if any(re.search(p, text, re.IGNORECASE) for p in DOCUMENT_PATTERNS['VOUCHER']):
                ws.cell(row=row, column=19, value="X")            # Col S

            if any(re.search(p, text, re.IGNORECASE) for p in DOCUMENT_PATTERNS['INFOGAS']):
                ws.cell(row=row, column=20, value="X")            # Col T

            if any(re.search(p, text, re.IGNORECASE) for p in DOCUMENT_PATTERNS['TUC']):
                ws.cell(row=row, column=21, value="X")            # Col U

            if any(re.search(p, text, re.IGNORECASE) for p in DOCUMENT_PATTERNS['RTV_ANTERIOR']):
                ws.cell(row=row, column=22, value="X")            # Col V

            if any(re.search(p, text, re.IGNORECASE) for p in DOCUMENT_PATTERNS['CONSULTA_CITV']):
                ws.cell(row=row, column=23, value="X")            # Col W

            if any(re.search(p, text, re.IGNORECASE) for p in DOCUMENT_PATTERNS['GLP_INICIAL_RENO']):
                ws.cell(row=row, column=25, value="X")            # Col Y

    out_path = os.path.join(app.config['UPLOAD_FOLDER'], "CHECKLIST_PROCESADO.xlsx")
    wb.save(out_path)
    return out_path

HTML_TEMPLATE = """
<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8">
    <title>Sistema de Verificación Automática CITV</title>
    <style>
        body { font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif; background-color: #f1f5f9; color: #1e293b; margin: 0; padding: 40px 20px; }
        .container { max-width: 750px; background: #ffffff; margin: 0 auto; padding: 35px; border-radius: 12px; box-shadow: 0 10px 25px rgba(0,0,0,0.05); }
        h1 { color: #0f172a; margin-top: 0; font-size: 26px; border-bottom: 3px solid #2563eb; padding-bottom: 12px; }
        .card { background: #f8fafc; border: 1px solid #e2e8f0; border-radius: 8px; padding: 20px; margin-top: 20px; }
        label { font-weight: 600; display: block; margin-bottom: 8px; color: #334155; }
        input[type="file"] { width: 100%; padding: 10px; border: 1px solid #cbd5e1; border-radius: 6px; background: #fff; box-sizing: border-box; }
        .btn { background: #2563eb; color: white; font-size: 16px; font-weight: 600; border: none; padding: 14px 20px; border-radius: 8px; cursor: pointer; width: 100%; margin-top: 25px; transition: background 0.2s; }
        .btn:hover { background: #1d4ed8; }
        .info { font-size: 13px; color: #64748b; margin-top: 6px; }
    </style>
</head>
<body>
    <div class="container">
        <h1>🚗 Sistema CITV - Verificación Automática</h1>
        <p>Sube tu archivo de Checklist en Excel y los documentos digitalizados para auditar de forma automática.</p>
        
        <form action="/upload" method="post" enctype="multipart/form-data">
            <div class="card">
                <label>1. Archivo Excel (Checklist Base):</label>
                <input type="file" name="excel_file" accept=".xlsx" required>
            </div>

            <div class="card">
                <label>2. Carpeta del Día (Contiene Subcarpetas por Placa y Archivos Sueltos):</label>
                <input type="file" name="doc_files" webkitdirectory directory multiple required>
                <div class="info">Selecciona la carpeta principal. El sistema procesará todas las subcarpetas y archivos sueltos internamente.</div>
            </div>

            <button type="submit" class="btn">⚡ Auditar Expedientes y Descargar Excel</button>
        </form>
    </div>
</body>
</html>
"""

@app.route('/')
def index():
    return render_template_string(HTML_TEMPLATE)

@app.route('/upload', methods=['POST'])
def upload():
    excel_file = request.files.get('excel_file')
    doc_files = request.files.getlist('doc_files')

    if not excel_file:
        return "Error: Falta el archivo Excel", 400

    excel_path = os.path.join(app.config['UPLOAD_FOLDER'], excel_file.filename)
    excel_file.save(excel_path)

    docs_dir = os.path.join(app.config['UPLOAD_FOLDER'], 'docs')
    os.makedirs(docs_dir, exist_ok=True)

    for f in doc_files:
        if f.filename:
            fpath = os.path.join(docs_dir, f.filename)
            os.makedirs(os.path.dirname(fpath), exist_ok=True)
            f.save(fpath)

    output_excel = process_expediente_excel(excel_path, docs_dir)
    return send_file(output_excel, as_attachment=True, download_name="CHECKLIST_PROCESADO.xlsx")

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000)
